"""Build the course report from the frozen, completed experiment records.

Run with the Codex bundled Python. Figures and numbers are actual experiment
outputs. This script does not train, evaluate, upload, or send anything.
"""
from pathlib import Path
import json
import hashlib
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
RUN = "20261027_134747_lego_main"
REPOSITORY_URL = "https://github.com/wwwcy9995/3dgs-view-count-study"
RECORDS = ROOT / "records/experiments" / RUN
OUT = ROOT / "output/report"
OUT.mkdir(parents=True, exist_ok=True)
DATA = json.loads((RECORDS / "summary.json").read_text(encoding="utf-8"))
SESSION = json.loads((RECORDS / "session.json").read_text(encoding="utf-8"))
assert SESSION["status"] == "complete"
ROWS = DATA["rows"]
assert [r["train_views"] for r in ROWS] == [10, 20, 40]
assert all(r["test_views"] == 200 and r["steps"] == 30000 for r in ROWS)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.85), Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(0.7), Cm(0.75)

def set_font(obj, east="宋体", latin="Times New Roman", size=11, bold=False):
    obj.font.name = latin
    obj.font.size = Pt(size)
    obj.font.bold = bold
    obj.font.color.rgb = RGBColor(0, 0, 0)
    rpr = obj.element.get_or_add_rPr() if hasattr(obj, "element") else obj._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + attr), latin)
    fonts.set(qn("w:eastAsia"), east)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)

for style_name in ("Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
    style = doc.styles[style_name]
    set_font(style)
    pf = style.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.18
    pf.widow_control = True
    ppr = style.element.get_or_add_pPr()
    grid = OxmlElement("w:snapToGrid"); grid.set(qn("w:val"), "0"); ppr.append(grid)

for name, size in (("Heading 1", 13), ("Heading 2", 11.5)):
    style = doc.styles[name]
    set_font(style, east="黑体", size=size, bold=True)
    style.paragraph_format.space_before = Pt(7)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

set_font(doc.styles["Title"], east="黑体", size=17, bold=True)
doc.styles["Title"].paragraph_format.line_spacing = 1.15
doc.styles["Title"].paragraph_format.space_after = Pt(9)
set_font(doc.styles["Caption"], size=9)
doc.styles["Caption"].paragraph_format.line_spacing = 1.12
doc.styles["Caption"].paragraph_format.space_after = Pt(5)

# Do not use the default Word line grid for Chinese body text.
grid = sec._sectPr.find(qn("w:docGrid"))
if grid is not None:
    sec._sectPr.remove(grid)

# The bundled base template carries blue title borders; remove them explicitly.
for el in list(doc.styles.element.iter(qn("w:pBdr"))):
    el.getparent().remove(el)

NEXT_PAGE = False

def para(text, *, size=None, bold=False, first=True, after=None, align=None, keep=False):
    p = doc.add_paragraph()
    previous = p._p.getprevious()
    if previous is not None and previous.tag == qn("w:tbl"):
        p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.first_line_indent = Pt(22) if first else Pt(0)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    if size is not None or bold:
        set_font(r, size=size or 11, bold=bold)
    return p

def heading(text, level=1):
    global NEXT_PAGE
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if NEXT_PAGE:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)
        NEXT_PAGE = False
    return p

def page():
    global NEXT_PAGE
    NEXT_PAGE = True

def caption(text):
    p = doc.add_paragraph(text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def figure(path, width=17.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))

def table(headers, body, widths, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths): column.width = Cm(width)
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        for k, v in {"val": "single", "sz": "5", "color": "D9D9D9"}.items():
            el.set(qn("w:" + k), v)
        borders.append(el)
    props.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", 65), ("bottom", 65), ("left", 85), ("right", 85)):
        el = OxmlElement("w:" + edge); el.set(qn("w:w"), str(value)); el.set(qn("w:type"), "dxa"); mar.append(el)
    props.append(mar)
    for values in [headers] + list(body):
        row = t.rows[0] if values is headers else t.add_row()
        rowpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); rowpr.append(cant)
        if values is headers:
            rep = OxmlElement("w:tblHeader"); rowpr.append(rep)
        for i, value in enumerate(values):
            cell = row.cells[i]; cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            set_font(p.add_run(str(value)), size=font_size, bold=(values is headers))
            if values is headers:
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "E8EEF3"); cell._tc.get_or_add_tcPr().append(shd)
    return t

def math_line(text):
    # A native editable OMML equation. No rasterized formulas or raw LaTeX.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    om = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), "Cambria Math"); fonts.set(qn("w:hAnsi"), "Cambria Math"); props.append(fonts)
    mr.append(props)
    txt = OxmlElement("m:t"); txt.text = text; mr.append(txt); om.append(mr); p._p.append(om)
    return p

# Page numbers are useful for a course report; no decorative header.
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run(""), size=9)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

doc.core_properties.title = "训练视角数量对 3D Gaussian Splatting 新视角合成质量的影响"
doc.core_properties.subject = "深度学习课程设计 2026 秋季"
doc.core_properties.author = ""
doc.core_properties.keywords = "3D Gaussian Splatting, Splatfacto, Lego, 训练视角数量, 新视角合成"

# PAGE 1: question, motivation and method.
p = doc.add_paragraph("训练视角数量对 3D Gaussian Splatting\n新视角合成质量的影响", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para("深度学习课程设计　2026 年秋季", first=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, after=5)
para("学号：________________　姓名：________________", first=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
link_paragraph = para("代码与实验记录：", first=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, after=7)
link = OxmlElement("w:hyperlink")
link.set(qn("r:id"), doc.part.relate_to(REPOSITORY_URL, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
link_run = OxmlElement("w:r")
link_properties = OxmlElement("w:rPr")
link_size = OxmlElement("w:sz"); link_size.set(qn("w:val"), "19"); link_properties.append(link_size)
link_fonts = OxmlElement("w:rFonts"); link_fonts.set(qn("w:ascii"), "Times New Roman"); link_fonts.set(qn("w:hAnsi"), "Times New Roman"); link_properties.append(link_fonts)
link_run.append(link_properties)
link_text = OxmlElement("w:t"); link_text.text = REPOSITORY_URL; link_run.append(link_text)
link.append(link_run); link_paragraph._p.append(link)
para("摘要", first=False, bold=True, after=3)
para("本文研究已知相机参数条件下，训练视角数量对 3D Gaussian Splatting 新视角合成质量的影响。基于 Nerfstudio 的 Splatfacto，在公开 NeRF Synthetic Lego 场景上构造 10、20、40 个视角的嵌套训练集，共用随机初始化点云，以相同分辨率和 30000 步预算训练，并在固定 200 个测试视角上评估。三组平均 PSNR 分别为 25.8787、30.3059、34.6534 dB，SSIM 和 LPIPS 也随视角增加而改善。对照图显示，稀疏组更容易出现镂空区域误填、凸点纹理模糊和边缘伪影。结论仅针对本次单场景、单种子实验，不代表每个测试视角都会改善。", size=10.5, first=False, after=5)
para("关键词：3D Gaussian Splatting；Splatfacto；训练视角数量；新视角合成", size=10, first=False, after=5)

heading("1 研究问题与相关工作")
para("新视角合成要求模型根据一组已知视角图像，生成未参与训练的相机视角图像。实际采集时，照片数量受时间和设备条件限制；减少训练视角会降低数据获取成本，也可能使遮挡背面、细杆和孔洞等区域缺少充分约束。因此，本文把问题限定为：同一场景在相同训练步数下，使用多少个已知位姿视角，会怎样影响测试图像质量，哪些局部区域最容易失真？")
para("3DGS 使用可优化的三维高斯表达场景，并通过可微光栅化拟合训练图像[1]。本文选择 Splatfacto，是因为其训练、渲染和评估接口便于形成完整实验流程。它属于 Nerfstudio 中的 3DGS 实现，底层使用 gsplat[2–4]，与原论文的 INRIA 训练代码并非同一实现。FSGS 等工作专门研究少视角合成[6]；本文只将其作为相关研究，不复现其方法，也不声称提出新算法。")

heading("2 方法原理")
para("每个三维高斯包含中心、尺度、旋转、不透明度和颜色等可学习参数。高斯经相机投影后在图像平面形成椭圆足迹，按深度进行透明度合成。Splatfacto 通过训练图与渲染图的光度误差反向传播，更新这些参数；颜色使用最高三阶球谐系数，以表达随观察方向变化的外观。训练过程中同时进行自适应增密和裁剪。")
para("本实验采用实现中的默认光度损失组合：平均绝对误差与结构相似性损失的权重分别为 0.8 和 0.2。主要方法设置保持一致，只改变训练图像数量。由此比较的是既定实现与固定优化预算下的整体效果，而非一种新网络或严格相同最终参数量下的效果。")

# PAGE 2: dataset, selection, controls.
page()
heading("3 数据与实验设计")
heading("3.1 公开数据与任务边界", 2)
para("数据采用 NeRF 作者公开的 Synthetic 数据集 Lego 场景[5]。原始数据包含 100 张训练图、100 张验证图、200 张测试图，均为 800×800 RGBA 合成图，并提供 Blender 相机内外参数。下载作者公开归档后只提取 Lego，训练和评估统一缩放为 400×400，将透明区域合成为白色背景。")
para("所有组直接使用公开的真实相机参数，不从少量照片重新估计位姿。因此，本研究回答的是“已知相机条件下的少视角外观学习”问题，不能表述为严格只输入 10、20 或 40 张照片的端到端三维重建。")
heading("3.2 嵌套训练子集与固定测试集", 2)
para("在原 100 个训练候选相机中心归一化后，以原索引 0 为起点，按方向最远点遍历依次选取相机；分别取序列前 10、20、40 项，形成嵌套子集，尽量覆盖相近观察范围。该选择使用全部候选训练相机的参数，但不使用测试图像的误差或质量指标。准确图像清单保存于 lego-splits.json。")
figure(ROOT / "records/data/camera-coverage.png", width=16.8)
caption("图 1　三组训练相机的覆盖情况　灰色点为固定测试相机")
para("三组均在原测试集的全部 200 张图上评估，测试图不参与训练或调参。训练与测试相机不重合；200 个测试相机位姿逐一核对为不同位姿。首次 1000 步试跑仅使用独立验证集中的 8 张图检查流程，不纳入本文正式结果。")
heading("3.3 初始化与预算控制", 2)
para("为避免全量图片生成的 SfM 点云提供额外几何信息，三组共用同一份与图像无关的随机点云：在 [−1.5, 1.5]³ 中均匀采样 10000 点，颜色统一为灰色，NumPy PCG64 种子为 42。该 PLY 文件及其 SHA-256 在训练前校验；各组从头训练，不加载试跑模型。")
table(["控制项", "统一设置"], [
    ("图像与测试集", "400×400　白背景　固定 200 张测试图"),
    ("优化预算", "每组 30000 步　随机种子 42　最终检查点"),
    ("初始化与相机", "同一随机点云　已知相机　关闭相机优化"),
    ("模型与执行", "三阶球谐　相同增密和裁剪设置　CPU 线程 4")
], [3.5, 13.5], font_size=9.5)
para("固定步数不等于每张图像获得相同更新次数，也不等于相同耗时或算力消耗；最终高斯数量由自适应训练决定，本文另外记录这一规模差异。", size=10, first=False, after=0)

# PAGE 3: implementation and evaluation.
page()
heading("4 设计实现与评价方法")
heading("4.1 实际运行环境", 2)
table(["项目", "本次实验环境"], [
    ("计算机", "Intel Core i7-12700H　内存约 15.7 GiB"),
    ("显卡与驱动", "RTX 3070 Ti Laptop GPU　8 GB　592.27"),
    ("系统", "Windows 11　WSL2　Ubuntu 22.04"),
    ("Python 与 PyTorch", "Python 3.10.22　PyTorch 2.1.2+cu118"),
    ("渲染与训练框架", "gsplat 1.4.0+pt21cu118　Nerfstudio 1.1.5"),
    ("其他关键依赖", "torchvision 0.16.2+cu118　NumPy 1.26.4")
], [3.5, 13.5], font_size=9.5)
para("训练运行在独立 Conda 环境 splat-course 中，使用 PyTorch 携带的 CUDA 11.8 运行时与匹配的 gsplat 预编译包。关闭 torch.compile，图像以 uint8 缓存到 CPU。日志中保留了 cuDNN 缺少 nvrtc 时的回退提示；三组在同一环境下均完成训练和评估。", after=6)
heading("4.2 项目实现与执行流程", 2)
para("项目在上游训练器之外实现了数据下载与校验、嵌套划分、随机点云生成、实验调度、逐图评价及绘图。调度脚本在启动前核对依赖版本、划分、初始化和主要配置，并将协议与执行代码快照写入运行目录。各组主要模型和优化配置的指纹一致。")
para("执行顺序为 10、20、40 视角分别训练；三组全部成功后，才在正式测试集上运行评估和导出。每组统一采用最终第 29999 步检查点，即已完成 30000 次迭代，不根据测试分数挑选最佳模型。训练、评估、导出共 9 个阶段全部正常退出。")
heading("4.3 图像指标与资源记录", 2)
para("评价脚本复用 Nerfstudio 模型的 PSNR、SSIM 和 LPIPS 实现，在量化为 PNG 之前的浮点 RGB 上计算。PSNR 依据归一化图像的均方误差计算，单位为 dB，越高越好；SSIM 衡量局部结构相似性，越高越好；LPIPS 使用预训练 AlexNet 特征衡量感知差异，越低越好。AlexNet 权重仅用于评价，并非预训练的三维场景模型。")
math_line("PSNR = −10 log₁₀(MSE)")
para("逐图计算后对 200 张测试图取算术平均，因此报告中的平均 PSNR 不是先合并所有图像 MSE 再换算的数值。评价脚本先在试跑验证图上与 ns-eval 结果核对一致。另记录训练进程单调计时、系统起止时刻跨度、导出 PLY 字节数和最终高斯数量；训练计时含进程启动及加载，不含评估与导出。")
heading("4.4 可追溯记录", 2)
para("正式运行编号为 20261027_134747_lego_main，上游提交为 6b60855003011b2ca23c2fe3f8e2ca6314c69924。复现入口见 README.md 中的 06_run_experiments.sh 与 07_analyze_experiments.sh，精确命令在 session.json。项目保留依赖、配置、种子、划分、日志、模型及其校验值，可由 summary.json 和 per-view.csv 重新汇总结果。", size=10.5)

# PAGE 4: actual quantitative results.
page()
heading("5 实验结果与定量分析")
para("表 1 汇总全部固定测试图的平均指标；图 2 展示相同数据的变化趋势。结果来自各组实际完成的正式训练，未使用试跑验证分数。", after=5)
caption("表 1　固定 200 个测试视角上的平均图像质量")
table(["训练视角", "PSNR / dB ↑", "SSIM ↑", "LPIPS ↓"], [
    (r["train_views"], f'{r["psnr_db"]:.4f}', f'{r["ssim"]:.6f}', f'{r["lpips"]:.6f}') for r in ROWS
], [2.7, 4.8, 4.8, 4.7], font_size=10)
figure(RECORDS / "figures/quality-vs-views.png", width=17.0)
caption("图 2　训练视角数量与测试图像质量　连线仅用于展示本次三个观测点")
para("从 10 增至 20 个视角，平均 PSNR 提升 4.4271 dB；从 20 增至 40 个视角，再提升 4.3476 dB。40 视角相对 10 视角累计提升 8.7747 dB，同时 SSIM 从 0.893821 提升至 0.979561，LPIPS 从 0.056021 降至 0.010686。在本次设置下，三个指标共同支持更多训练视角带来更好的平均新视角质量。")
para("平均值不能代表每一张图：40 相比 20 视角的 PSNR 在 191/200 张图上更高，仍有 9 张更低；40 相比 10 视角也有 2 张 PSNR 更低。三组的逐图 PSNR 标准差分别为 2.5947、2.7993、3.4866 dB，只反映测试视角之间的难度差异，不是重复训练的不确定性或置信区间。")
caption("表 2　训练时间记录与导出模型规模")
table(["视角数", "单调计时 / s", "起止跨度 / s", "PLY / MiB", "高斯数量"], [
    (r["train_views"], f'{r["train_process_elapsed_seconds"]:.2f}', f'{r["start_finish_clock_seconds"]:.2f}', f'{r["ply_mib"]:.3f}', str(r["gaussians"])) for r in ROWS
], [2.0, 3.9, 3.9, 3.4, 3.8], font_size=9.5)
para("20 视角组的单调计时为 514.17 s，而系统起止时刻跨度为 1493.98 s，两者相差约 979.81 s，具体原因未确定。本文如实保留两类计时，不据此对三组作严格速度排序。MiB 按 2²⁰ 字节计算，模型大小统一采用导出高斯 PLY，训练检查点大小另行归档。", size=10.5)
para("40 视角组的高斯数量和 PLY 大小略低于另外两组，但质量更高，说明此次改善并不伴随最终高斯数量的简单增加。由于增密和裁剪是自适应过程，本文仍不能将其解释为严格控制模型容量后的单一因素因果结论。", size=10.5)

# PAGE 5: predetermined qualitative cases, unchanged output figures.
page()
heading("6 同视角渲染与局部分析")
para("为避免按渲染结果挑图，协议预先固定测试索引 0、50、100、150；本文展示其中 0 和 100。每幅图按“真实图、10、20、40 视角”排列，上行为全图，下行为同坐标局部放大。裁剪中心取自真实图前景包围盒，宽和高各为包围盒的 40%，各组一致。", size=10.5, after=6)
figure(RECORDS / "figures/test-view-0000.png", width=16.6)
caption("图 3　测试视角 0　10 视角组在中央镂空处出现错误填色，凸点和平台纹理较模糊；40 视角更接近真实图")
figure(RECORDS / "figures/test-view-0100.png", width=16.6)
caption("图 4　测试视角 100　少视角组在窗框与细杆周围出现误填和边缘拖影，增加视角后有所减轻")
para("这些区域涉及细小结构、遮挡边界或孔洞；多角度观测可能提供更充分的约束。不过，仅由 RGB 渲染图可确认的是可见外观失真，不能证明伪影的三维位置或唯一成因。本文未计算几何误差，也不将彩色漂浮状斑块直接等同于已验证的几何错误。", size=10, after=0)

# PAGE 6: interpretation, limitations and proper attribution.
page()
heading("7 分析思考与结论")
heading("7.1 结果解释与局限", 2)
para("定量指标与同视角对照给出一致证据：在本次 Lego 场景中，增加训练视角后，平均图像质量提升，镂空区域误填和局部细节模糊减轻。一种合理解释是，更多方向的观测使颜色和可见性获得更充分约束；但本文没有额外消融实验来分离视角数量、空间覆盖和优化过程各自的作用。")
para("实验只有一个合成场景和一个训练种子，不能推广到真实拍摄、复杂背景或全部随机初始化。已知真实相机参数降低了任务难度；三组又采用覆盖较广的嵌套子集，因此结果不代表相同数量但集中在一侧的照片。白色背景在全图指标中占较大比例，局部细节仍需结合放大图分析。")
para("固定 30000 步使预算口径清楚，但训练图数量增加时，每张图平均被采样的机会会变化，尚未验证各组是否同等接近收敛。自适应增密与裁剪带来不同最终规模；此外，单次运行的时间存在时钟差异，不能用于严谨的效率结论。后续可增加训练种子和第二场景，并对固定数量下均匀与集中视角、固定步数与更充分收敛分别作比较。")
heading("7.2 结论", 2)
para("在已知相机、共同随机点云初始化、400×400 和相同 30000 步条件下，本次 Splatfacto 的 10、20、40 训练视角实验得到逐步提高的平均测试质量。40 视角相对 10 视角平均 PSNR 提升 8.7747 dB，局部外观更接近真实图；部分单张测试图仍有退步。该结论是一次受控的小规模实验观察，不是对任意场景的普遍保证。")
heading("8 代码来源与工具使用说明")
para("场景数据来自 NeRF 作者公开数据；训练方法基于 Nerfstudio Splatfacto 和 gsplat 开源实现，按上游许可证使用并记录版本，不将其称为自行提出的算法。项目新增数据划分、共同初始化、实验调度、指标汇总及绘图脚本。使用 OpenAI Codex 协助环境排错、脚本编写、结果整理与文字编辑；文中指标及图像均来自本机实际运行。项目代码与实验记录整理为配套仓库，地址见首页。", size=10.5)

heading("参考文献")
refs = [
    "[1] Kerbl B, Kopanas G, Leimkühler T, et al. 3D Gaussian Splatting for Real-Time Radiance Field Rendering. ACM Transactions on Graphics, 2023. https://repo-sam.inria.fr/fungraph/3d-gaussian-splatting/",
    "[2] Nerfstudio. Splatfacto documentation. https://docs.nerf.studio/nerfology/methods/splat.html （访问日期 2026-10-27）.",
    "[3] Tancik M, Weber E, Ng E, et al. Nerfstudio: A Modular Framework for Neural Radiance Field Development. 2023. arXiv:2302.04264.",
    "[4] Ye V, Li R, Kerr J, et al. gsplat: An Open-Source Library for Gaussian Splatting. 2024. arXiv:2409.06765.",
    "[5] Mildenhall B, Srinivasan P P, Tancik M, et al. NeRF: Representing Scenes as Neural Radiance Fields for View Synthesis. ECCV, 2020. 数据与下载脚本：https://github.com/bmild/nerf .",
    "[6] Zhu Z, Fan Z, Jiang Y, et al. FSGS: Real-Time Few-shot View Synthesis using Gaussian Splatting. 2023. arXiv:2312.00451."
]
for text in refs:
    p = para(text, first=False, size=9, after=4)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.left_indent = Pt(14)
    p.paragraph_format.first_line_indent = Pt(-14)

target = OUT / "训练视角数量与3DGS新视角合成_课程报告.docx"
doc.save(target)
provenance = {
    "report": str(target.relative_to(ROOT)),
    "experiment_run": RUN,
    "sources": {},
    "requested_format": "docx",
    "student_name_and_id": "left blank as requested",
    "created_from_real_results": True,
    "repository_url": REPOSITORY_URL,
}
for path in [RECORDS / "summary.json", RECORDS / "session.json", ROOT / "records/data/camera-coverage.png",
             RECORDS / "figures/quality-vs-views.png", RECORDS / "figures/test-view-0000.png", RECORDS / "figures/test-view-0100.png"]:
    provenance["sources"][str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
provenance["report_sha256"] = hashlib.sha256(target.read_bytes()).hexdigest()
(OUT / "report-provenance.json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
print(target)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")
